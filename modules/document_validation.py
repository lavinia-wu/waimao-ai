from __future__ import annotations

from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
import json
import re
import unicodedata
import zipfile
from xml.etree.ElementTree import ParseError
from typing import Any

from database.db import get_connection
from modules.documents import DOCUMENT_TYPES, get_order_document_context
from modules.lc_deadlines import parse_date


FIELD_DEFINITIONS = [
    ('goods_description', '货物描述'),
    ('product_name', '品名'),
    ('quantity', '数量'),
    ('unit_price', '单价'),
    ('total_amount', '总金额'),
    ('currency', '币别'),
    ('incoterm', '贸易条款'),
    ('loading_port', '起运港'),
    ('destination_port', '目的港'),
    ('latest_shipment_date', '最迟装运日'),
    ('presentation_deadline', '交单截止日'),
    ('lc_expiry_date', '信用证有效期'),
    ('gross_weight', '毛重'),
    ('net_weight', '净重'),
    ('volume', '体积'),
    ('marks_and_numbers', '唛头'),
    ('document_counts', '单据份数'),
]

WORD_ARTIFACT_FIELDS = {
    'PI': {
        'goods_description', 'product_name', 'quantity', 'unit_price',
        'total_amount', 'currency', 'incoterm', 'loading_port',
        'destination_port', 'latest_shipment_date',
    },
    'Commercial Invoice': {
        'goods_description', 'product_name', 'quantity', 'unit_price',
        'total_amount', 'currency', 'incoterm', 'loading_port',
        'destination_port', 'latest_shipment_date', 'gross_weight',
        'net_weight', 'volume', 'marks_and_numbers',
    },
    'Packing List': {
        'goods_description', 'product_name', 'quantity', 'loading_port',
        'destination_port', 'gross_weight', 'net_weight', 'volume',
        'marks_and_numbers',
    },
    'Sales Contract': {
        'goods_description', 'product_name', 'quantity', 'unit_price',
        'total_amount', 'currency', 'incoterm', 'loading_port',
        'destination_port', 'latest_shipment_date', 'gross_weight',
        'net_weight', 'volume', 'marks_and_numbers',
    },
}

EXCEL_ARTIFACT_FIELDS = {
    'Commercial Invoice': {
        'goods_description', 'product_name', 'quantity', 'unit_price',
        'total_amount', 'currency', 'incoterm', 'loading_port',
        'destination_port', 'gross_weight', 'net_weight',
    },
    'Packing List': {
        'product_name', 'quantity', 'gross_weight', 'net_weight',
    },
}

NUMERIC_FIELDS = {
    'quantity', 'unit_price', 'total_amount', 'gross_weight', 'net_weight', 'volume',
}
DATE_FIELDS = {'latest_shipment_date', 'presentation_deadline', 'lc_expiry_date'}
REQUIRED_FIELDS = {
    'goods_description', 'product_name', 'quantity', 'total_amount', 'currency',
    'incoterm', 'loading_port', 'destination_port',
}
FIELD_SEVERITY = {field: 'high' for field, _ in FIELD_DEFINITIONS}
FIELD_SEVERITY['marks_and_numbers'] = 'medium'
ARTIFACT_LABELS = {
    'goods_description': ('description of goods', 'goods description', 'product description', 'description', '货物描述', '货描'),
    'product_name': ('product name', 'commodity', '品名', '产品名称'),
    'quantity': ('quantity', 'qty', '数量'),
    'unit_price': ('unit price', 'price per unit', 'unit value', '单价'),
    'total_amount': ('total amount', 'total value', 'amount', '总金额', '总价'),
    'currency': ('currency', '币种', '币别'),
    'incoterm': ('incoterm', 'trade term', 'terms of delivery', '贸易条款'),
    'loading_port': ('port of loading', 'loading port', '起运港', '装运港'),
    'destination_port': ('port of discharge', 'port of destination', 'destination port', '目的港', '卸货港'),
    'latest_shipment_date': ('latest date of shipment', 'latest shipment date', 'shipment date', '最迟装运日', '装运日期'),
    'presentation_deadline': ('presentation deadline', 'latest date for presentation', '交单截止日'),
    'lc_expiry_date': ('expiry date', 'expiration date', '信用证有效期'),
    'gross_weight': ('gross weight', 'g.w.', '毛重'),
    'net_weight': ('net weight', 'n.w.', '净重'),
    'volume': ('volume', 'measurement', 'cbm', '体积'),
    'marks_and_numbers': ('marks and numbers', 'shipping marks', 'marks', '唛头'),
}
PORT_ALIASES = {
    'shanghai': 'shanghaichina',
    'shanghaichina': 'shanghaichina',
    'jebelali': 'jebelalidubai',
    'jebelalidubai': 'jebelalidubai',
    'rotterdam': 'rotterdamnetherlands',
    'rotterdamnetherlands': 'rotterdamnetherlands',
}


def _json_object(value: Any, default: Any) -> Any:
    try:
        return json.loads(value) if value else default
    except (TypeError, json.JSONDecodeError):
        return default


def _decimal(value: Any) -> Decimal | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float, Decimal)):
        try:
            result = Decimal(str(value))
            return result if result.is_finite() else None
        except InvalidOperation:
            return None
    text = str(value).strip()
    if not text:
        return None
    negative = text.startswith('(') and text.endswith(')')
    text = text.strip('()').strip()
    text = re.sub(r'(?i)\b(?:US|CA|AU)\$', '', text)
    text = re.sub(r'(?i)\b(?:USD|EUR|GBP|CNY|RMB|JPY|CAD|AUD)\b', '', text)
    text = re.sub(r'[$€£¥￥]', '', text).strip()
    text = re.sub(r'(?i)\s*(?:MT|KG|KGS|TONS?|CBM|M3|PCS|SETS?)\s*$', '', text)
    match = re.fullmatch(r'[-+]?\d[\d,]*(?:\.\d+)?', text)
    if not match:
        return None
    number = text
    if ',' in number:
        if re.fullmatch(r'[-+]?\d{1,3}(?:,\d{3})+(?:\.\d+)?', number):
            number = number.replace(',', '')
        else:
            return None
    try:
        result = Decimal(number)
    except InvalidOperation:
        return None
    if not result.is_finite():
        return None
    return -abs(result) if negative else result


def _number(value: Any) -> float | None:
    parsed = _decimal(value)
    return float(parsed) if parsed is not None else None


def _normal_text(value: Any) -> str:
    text = unicodedata.normalize('NFKC', str(value or '')).casefold()
    return ''.join(
        character for character in text
        if character.isalnum()
    )


def _normalized_port(value: Any) -> str:
    normalized = _normal_text(value)
    return PORT_ALIASES.get(normalized, normalized)


def _compare_values(
    field: str,
    left: Any,
    right: Any,
    core_product_name: Any = None,
) -> bool | None:
    if left in (None, '') or right in (None, ''):
        return None
    if field in DATE_FIELDS:
        left_date, right_date = parse_date(left), parse_date(right)
        if isinstance(left, (date, datetime)):
            left_date = left.date() if isinstance(left, datetime) else left
        if isinstance(right, (date, datetime)):
            right_date = right.date() if isinstance(right, datetime) else right
        return left_date == right_date if left_date and right_date else None
    if field in NUMERIC_FIELDS:
        left_number, right_number = _decimal(left), _decimal(right)
        if left_number is None or right_number is None:
            return None
        tolerances = {
            'quantity': max(abs(left_number), abs(right_number)) * Decimal('0.001'),
            'unit_price': Decimal('0.01'),
            'total_amount': Decimal('0.01'),
            'gross_weight': max(abs(left_number), abs(right_number)) * Decimal('0.005'),
            'net_weight': max(abs(left_number), abs(right_number)) * Decimal('0.005'),
            'volume': max(abs(left_number), abs(right_number)) * Decimal('0.005'),
        }
        return abs(left_number - right_number) <= tolerances[field]
    left_text = _normalized_port(left) if field in {'loading_port', 'destination_port'} else _normal_text(left)
    right_text = _normalized_port(right) if field in {'loading_port', 'destination_port'} else _normal_text(right)
    if not left_text or not right_text:
        return None
    if left_text == right_text:
        return True
    if field == 'goods_description':
        core_name = _normal_text(core_product_name)
        if not core_name:
            return None
        return core_name in left_text and core_name in right_text
    return False


def _values_equal(field: str, left: Any, right: Any) -> bool | None:
    return _compare_values(field, left, right)


def _is_blocking_issue(issue: dict[str, Any]) -> bool:
    return issue.get('status') == 'error' and issue.get('severity') == 'high'


def _unresolved_high_risk_blockers(risks: list[dict[str, Any]]) -> list[dict[str, Any]]:
    blockers = []
    for risk in risks:
        if (
            str(risk.get('risk_level') or '').strip().casefold() in {'高', 'high', 'critical'}
            and risk.get('handling_status') not in {'已确认可接受', '已整改'}
        ):
            blockers.append({
                'field': f"信用证高风险条款 · {risk.get('risk_type') or risk.get('id')}",
                'status': 'error',
                'severity': 'high',
                'reasons': [
                    f"处理状态为“{risk.get('handling_status') or '未处理'}”，"
                    '请先在信用证风控页标记为“已确认可接受”或“已整改”。'
                ],
            })
    return blockers


def _lc_values(lc: dict[str, Any] | None, parsed: dict[str, Any]) -> dict[str, Any]:
    lc = lc or {}
    return {
        'goods_description': parsed.get('goods_description') or lc.get('goods_description'),
        'product_name': parsed.get('product_name_en') or parsed.get('product_name'),
        'quantity': parsed.get('quantity'),
        'unit_price': parsed.get('unit_price'),
        'total_amount': parsed.get('lc_amount') or lc.get('lc_amount'),
        'currency': parsed.get('currency') or lc.get('currency'),
        'incoterm': parsed.get('incoterm'),
        'loading_port': parsed.get('loading_port') or lc.get('loading_port'),
        'destination_port': parsed.get('discharge_port') or lc.get('discharge_port'),
        'latest_shipment_date': parsed.get('latest_shipment_date') or lc.get('latest_shipment_date'),
        'presentation_deadline': parsed.get('presentation_deadline'),
        'lc_expiry_date': parsed.get('expiry_date') or lc.get('expiry_date'),
        'gross_weight': parsed.get('gross_weight'),
        'net_weight': parsed.get('net_weight'),
        'volume': parsed.get('volume'),
        'marks_and_numbers': parsed.get('marks_and_numbers'),
        'document_counts': None,
    }


def _order_values(data: dict[str, Any]) -> dict[str, Any]:
    return {
        'goods_description': data.get('product_description_en') or data.get('product_specification'),
        'product_name': data.get('product_name_en') or data.get('product_name_cn'),
        'quantity': data.get('quantity'),
        'unit_price': data.get('unit_price'),
        'total_amount': data.get('total_amount'),
        'currency': data.get('currency'),
        'incoterm': data.get('incoterm'),
        'loading_port': data.get('loading_port'),
        'destination_port': data.get('destination_port'),
        'latest_shipment_date': data.get('latest_shipment_date'),
        'presentation_deadline': data.get('presentation_deadline'),
        'lc_expiry_date': data.get('lc_expiry_date'),
        'gross_weight': data.get('gross_weight'),
        'net_weight': data.get('net_weight'),
        'volume': data.get('volume'),
        'marks_and_numbers': data.get('marks_and_numbers'),
        'document_counts': None,
    }


def _document_values(data: dict[str, Any]) -> dict[str, Any]:
    return {
        'goods_description': data.get('product_description_en') or data.get('product_specification'),
        'product_name': data.get('product_name_en') or data.get('product_name_cn'),
        'product_name_en': data.get('product_name_en') or data.get('product_name'),
        'product_specification': data.get('product_specification'),
        'quantity': data.get('quantity'),
        'unit_price': data.get('unit_price'),
        'total_amount': data.get('total_amount'),
        'currency': data.get('currency'),
        'incoterm': data.get('incoterm'),
        'loading_port': data.get('loading_port'),
        'destination_port': data.get('destination_port'),
        'latest_shipment_date': data.get('latest_shipment_date'),
        'presentation_deadline': data.get('presentation_deadline'),
        'lc_expiry_date': data.get('lc_expiry_date'),
        'gross_weight': data.get('gross_weight'),
        'net_weight': data.get('net_weight'),
        'volume': data.get('volume'),
        'marks_and_numbers': data.get('marks_and_numbers'),
        'document_counts': None,
    }


def _artifact_values(path: str) -> dict[str, Any]:
    """Extract only label-adjacent values; never infer fields from whole-document text."""
    from pathlib import Path

    source = Path(path)
    rows: list[list[str]] = []
    if source.suffix.lower() == '.docx':
        from docx import Document
        from docx.opc.exceptions import PackageNotFoundError

        try:
            document = Document(source)
        except (PackageNotFoundError, zipfile.BadZipFile, KeyError, ValueError, ParseError) as exc:
            raise ValueError(f'无法解析 Word 模板：{exc}') from exc
        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                rows.append([paragraph.text.strip()])
        for table in document.tables:
            rows.extend([[cell.text.strip() for cell in row.cells] for row in table.rows])
        for section in document.sections:
            for paragraph in (*section.header.paragraphs, *section.footer.paragraphs):
                if paragraph.text.strip():
                    rows.append([paragraph.text.strip()])
    elif source.suffix.lower() == '.xlsx':
        from openpyxl import load_workbook
        from openpyxl.utils.exceptions import InvalidFileException

        try:
            workbook = load_workbook(source, read_only=True, data_only=True)
        except (InvalidFileException, zipfile.BadZipFile, EOFError, KeyError, ValueError, ParseError) as exc:
            raise ValueError(f'无法解析 Excel 模板：{exc}') from exc
        try:
            for sheet in workbook.worksheets:
                for row in sheet.iter_rows():
                    rows.append([
                        str(cell.value).strip() if cell.value is not None else ''
                        for cell in row
                    ])
        finally:
            workbook.close()
    else:
        raise ValueError(f'不支持校验的单证文件格式：{source.suffix}')

    label_lookup = {
        _normal_text(alias): field
        for field, aliases in ARTIFACT_LABELS.items()
        for alias in aliases
    }
    extracted: dict[str, Any] = {}

    def assign(label: str, value: Any) -> None:
        normalized_label = _normal_text(label)
        field = label_lookup.get(normalized_label)
        if field is None:
            for alias, candidate in label_lookup.items():
                if normalized_label.startswith(alias) and normalized_label[len(alias):] in {
                    'mt', 'kg', 'cbm', 'usd', 'eur',
                }:
                    field = candidate
                    break
        if field and value not in (None, '') and not extracted.get(field):
            extracted[field] = value

    for row in rows:
        for index, cell in enumerate(row):
            text = str(cell or '').strip()
            if not text:
                continue
            inline = re.match(r'^(.{2,60}?)\s*[:：]\s*(.+)$', text)
            if inline:
                assign(inline.group(1), inline.group(2).strip())
                continue
            if index + 1 < len(row):
                assign(text, row[index + 1])
    return extracted


def _document_count_requirements(lc: dict[str, Any] | None, parsed: dict[str, Any]) -> list[dict[str, Any]]:
    raw = parsed.get('documents_required')
    if not isinstance(raw, list):
        raw = _json_object((lc or {}).get('documents_required_json'), [])
    if not isinstance(raw, list):
        raw = []
    result = []
    for index, item in enumerate(raw):
        if not isinstance(item, dict):
            continue
        original = _non_negative_count(item.get('original_count'))
        copies = _non_negative_count(item.get('copy_count'))
        if original is None and copies is None:
            continue
        result.append({
            'key': str(index),
            'lc_id': (lc or {}).get('id'),
            'document_type': str(item.get('document_type') or '单据'),
            'original_count': original or 0,
            'copy_count': copies or 0,
            'required_total': (original or 0) + (copies or 0),
            'requirements': item.get('special_requirements') or item.get('requirements') or '',
        })
    return result


def _non_negative_count(value: Any) -> int | None:
    if value is None or value == '':
        return None
    try:
        result = int(value)
    except (TypeError, ValueError):
        return None
    return result if result >= 0 else None


def get_document_count_requirements(order_id: int) -> list[dict[str, Any]]:
    with get_connection() as conn:
        lc_row = conn.execute(
            'SELECT * FROM letters_of_credit WHERE order_id = ? AND is_approved = 1 ORDER BY id DESC LIMIT 1',
            (order_id,),
        ).fetchone()
    lc = dict(lc_row) if lc_row else None
    parsed = _json_object((lc or {}).get('parsed_json'), {})
    return _document_count_requirements(lc, parsed)


def _count_status(
    requirements: list[dict[str, Any]],
    prepared_counts: dict[str, dict[str, int]] | None,
) -> tuple[str, list[str]]:
    if not requirements:
        return 'needs_review', []
    prepared_counts = prepared_counts or {}
    issues = []
    for item in requirements:
        counts = prepared_counts.get(item['key'], {})
        originals = _non_negative_count(counts.get('original_count')) or 0
        copies = _non_negative_count(counts.get('copy_count')) or 0
        if originals < item['original_count'] or copies < item['copy_count']:
            issues.append(
                f"{item['document_type']}：要求 {item['original_count']} 份正本、"
                f"{item['copy_count']} 份副本；确认备妥 {originals} 份正本、{copies} 份副本"
            )
    return ('error' if issues else 'passed', issues)


def validate_order_consistency(
    order_id: int,
    prepared_counts: dict[str, dict[str, int]] | None = None,
) -> dict[str, Any]:
    order_data = get_order_document_context(order_id)
    with get_connection() as conn:
        lc_row = conn.execute(
            'SELECT * FROM letters_of_credit WHERE order_id = ? AND is_approved = 1 ORDER BY id DESC LIMIT 1',
            (order_id,),
        ).fetchone()
        document_rows = conn.execute(
            'SELECT * FROM documents WHERE order_id = ? ORDER BY id DESC',
            (order_id,),
        ).fetchall()

    lc = dict(lc_row) if lc_row else None
    parsed_lc = _json_object((lc or {}).get('parsed_json'), {})
    lc_values = _lc_values(lc, parsed_lc)
    order_values = _order_values(order_data)
    calculated_deadline = computed_presentation_deadline(parsed_lc, order_data)
    if calculated_deadline:
        lc_values['presentation_deadline'] = calculated_deadline.isoformat()
    tolerance_percent = _number(parsed_lc.get('tolerance_percentage')) or 0
    documents: dict[str, dict[str, Any]] = {}
    for row in document_rows:
        item = dict(row)
        document_type = item.get('document_type')
        if document_type in DOCUMENT_TYPES and document_type not in documents:
            snapshot = _json_object(item.get('content_json'), {})
            source_path = item.get('file_path_word') or item.get('file_path_excel')
            artifact_values = None
            artifact_error = ''
            if source_path:
                try:
                    artifact_values = _artifact_values(source_path)
                except (OSError, ValueError, ImportError, KeyError) as exc:
                    artifact_error = str(exc)
            documents[document_type] = {
                'record': item,
                'values': _document_values(snapshot) if snapshot else {},
                'artifact_values': artifact_values,
                'artifact_error': artifact_error,
            }

    rows = []
    count_requirements = _document_count_requirements(lc, parsed_lc)
    needs_review = []
    warnings = []
    for field, label in FIELD_DEFINITIONS:
        values = {'信用证': lc_values.get(field), '订单': order_values.get(field)}
        for document_type in DOCUMENT_TYPES:
            document = documents.get(document_type)
            if not document:
                values[document_type] = None
                continue
            values[document_type] = (
                document['artifact_values'].get(field)
                if document['artifact_values'] and field in document['artifact_values']
                else document['values'].get(field)
            )

        known = [(source, value) for source, value in values.items() if value not in (None, '')]
        mismatches: list[str] = []
        uncertain: list[str] = []
        soft_differences: list[str] = []
        if len(known) > 1:
            baseline_source, baseline_value = known[0]
            for source, value in known[1:]:
                equal = _compare_values(
                    field,
                    baseline_value,
                    value,
                    core_product_name=order_data.get('product_name_en'),
                )
                if (
                    equal is False
                    and field == 'total_amount'
                    and (baseline_source == '信用证' or source == '信用证')
                ):
                    baseline_number = _decimal(baseline_value)
                    value_number = _decimal(value)
                    lc_amount = _decimal(lc_values.get('total_amount'))
                    allowed_difference = max(
                        Decimal('0.01'),
                        abs(lc_amount or Decimal(0)) * Decimal(str(tolerance_percent)) / 100,
                    )
                    equal = (
                        baseline_number is not None
                        and value_number is not None
                        and abs(baseline_number - value_number) <= allowed_difference
                    )
                    if equal:
                        soft_differences.append(
                            f'{baseline_source}={baseline_value}；{source}={value}，'
                            '差异在信用证允许的金额容差内，请人工确认'
                        )
                if equal is False:
                    mismatches.append(
                        f'{baseline_source}={baseline_value}；{source}={value}'
                    )
                elif equal is None:
                    uncertain.append(
                        f'{baseline_source}={baseline_value}；{source}={value} 无法可靠标准化比较'
                    )

        if field == 'document_counts':
            count_status, count_issues = _count_status(
                count_requirements,
                prepared_counts,
            )
            values['信用证'] = {
                'passed': '份数已确认',
                'error': '份数不足',
                'needs_review': '待核对',
            }.get(count_status, count_status)
            values['订单'] = '人工确认份数'
            mismatches = count_issues
            if not count_requirements:
                uncertain.append('未识别到信用证份数要求，需人工核对信用证原文')

        artifact_missing: list[str] = []
        artifact_unknown: list[str] = []
        artifact_mismatches: list[str] = []
        if field != 'document_counts':
            for document_type, document in documents.items():
                suffix = str(
                    document['record'].get('file_path_word')
                    or document['record'].get('file_path_excel')
                    or ''
                ).lower()
                if suffix.endswith('.docx'):
                    applicable_fields = WORD_ARTIFACT_FIELDS.get(document_type, set())
                elif suffix.endswith('.xlsx'):
                    applicable_fields = EXCEL_ARTIFACT_FIELDS.get(document_type, set())
                else:
                    applicable_fields = (
                        WORD_ARTIFACT_FIELDS.get(document_type, set())
                        | EXCEL_ARTIFACT_FIELDS.get(document_type, set())
                    )
                if field not in applicable_fields:
                    continue
                expected = document['values'].get(field)
                if expected in (None, ''):
                    artifact_missing.append(f'{document_type}：生成数据中缺少{label}')
                    continue
                artifact_values = document['artifact_values']
                actual = artifact_values.get(field) if artifact_values else None
                if actual in (None, ''):
                    reason = document['artifact_error'] or '模板中未找到可可靠关联的字段标签和值'
                    artifact_unknown.append(f'{document_type}：{reason}')
                    continue
                artifact_equal = _compare_values(
                    field,
                    expected,
                    actual,
                    core_product_name=order_data.get('product_name_en'),
                )
                if artifact_equal is False:
                    artifact_mismatches.append(
                        f'{document_type}文件提取值={actual}；生成数据={expected}'
                    )
                elif artifact_equal is None:
                    artifact_unknown.append(
                        f'{document_type}：文件值“{actual}”无法可靠标准化比较'
                    )

        mismatches.extend(artifact_mismatches)
        uncertain.extend(artifact_unknown)
        severity = FIELD_SEVERITY.get(field, 'medium')
        if field == 'document_counts' and mismatches:
            status = 'error'
            details = mismatches
        elif field == 'document_counts' and count_requirements:
            status = 'passed'
            details = []
        elif field == 'document_counts':
            status = 'needs_review'
            details = uncertain
        elif mismatches:
            status = 'error'
            details = mismatches
        elif artifact_missing:
            status = 'missing'
            details = artifact_missing
        elif uncertain or (known and len(known) < 2):
            status = 'needs_review'
            details = uncertain or ['可比较来源不足，需人工核对']
        elif soft_differences:
            status = 'warning'
            details = soft_differences
        elif not known:
            status = 'missing' if field in REQUIRED_FIELDS else 'needs_review'
            details = ['必填字段未提供' if status == 'missing' else '没有可比较的数据']
        else:
            status = 'passed'
            details = []

        row = {
            'field': field,
            'label': label,
            'values': values,
            'status': status,
            'severity': severity,
            'details': details,
        }
        rows.append(row)
        if status == 'needs_review':
            needs_review.append({'field': label, 'reasons': details})
        elif status == 'missing':
            warnings.append({'field': label, 'status': status, 'reasons': details})
        elif status == 'warning':
            warnings.append({'field': label, 'status': status, 'reasons': details})

    missing_documents = [doc_type for doc_type in DOCUMENT_TYPES if doc_type not in documents]
    for document_type in missing_documents:
        needs_review.append({
            'field': document_type,
            'reasons': [f'{document_type} 尚未生成，尚无文件内容可供校验。'],
        })

    blockers = [
        {'field': item['label'], 'status': item['status'], 'severity': item['severity'], 'reasons': item['details']}
        for item in rows
        if item['status'] == 'error' and item['severity'] == 'high'
    ]
    if lc:
        from modules.lc_parser import get_lc_risks

        blockers.extend(_unresolved_high_risk_blockers(get_lc_risks(int(lc['id']))))

    return {
        'order_id': order_id,
        'lc_id': lc.get('id') if lc else None,
        'rows': rows,
        'blockers': blockers,
        'needs_review': needs_review,
        'warnings': warnings,
        'missing_documents': missing_documents,
        'document_count_requirements': count_requirements,
        'has_lc': lc is not None,
    }


def assert_pdf_consistency(
    order_id: int,
    prepared_counts: dict[str, dict[str, int]] | None = None,
) -> dict[str, Any]:
    result = validate_order_consistency(order_id, prepared_counts)
    blocking_issues = [item for item in result['blockers'] if _is_blocking_issue(item)]
    if blocking_issues:
        messages = [
            f"{item['field']}：{';'.join(item['reasons'])}"
            for item in blocking_issues
        ]
        raise ValueError(
            '一致性校验未通过，已阻止生成 PDF。请先处理以下问题：\n- '
            + '\n- '.join(messages)
        )
    return result


def computed_presentation_deadline(parsed_lc: dict[str, Any], order: dict[str, Any]) -> date | None:
    explicit = parse_date(parsed_lc.get('presentation_deadline') or order.get('presentation_deadline'))
    if explicit:
        return explicit
    shipment = parse_date(order.get('shipment_date'))
    if not shipment:
        return None
    days = parsed_lc.get('presentation_period_days')
    try:
        days = int(days)
    except (TypeError, ValueError):
        match = re.search(r'(\d+)\s*days?', str(parsed_lc.get('presentation_period') or ''), re.I)
        days = int(match.group(1)) if match else None
    if days is None or days < 0:
        return None
    deadline = shipment + timedelta(days=days)
    expiry = parse_date(parsed_lc.get('expiry_date') or order.get('lc_expiry_date'))
    return min(deadline, expiry) if expiry else deadline
