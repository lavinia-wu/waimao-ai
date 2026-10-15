from __future__ import annotations

import json
import re
from io import BytesIO
from pathlib import Path
from typing import Any

from docx import Document
from openpyxl import load_workbook

from modules.lc_parser import _call_chat_completion, extract_lc_text


SUPPORTED_ORDER_EXTENSIONS = {
    '.pdf', '.doc', '.docx', '.xls', '.xlsx', '.xlsm', '.png', '.jpg', '.jpeg',
}


def _extract_legacy_word_text(content: bytes) -> str:
    import tempfile

    import win32com.client

    with tempfile.NamedTemporaryFile(suffix='.doc', delete=False) as source:
        source.write(content)
        path = Path(source.name)
    application = None
    try:
        application = win32com.client.DispatchEx('Word.Application')
        application.Visible = False
        application.DisplayAlerts = 0
        document = application.Documents.Open(str(path.resolve()), ReadOnly=True, AddToRecentFiles=False)
        try:
            return str(document.Content.Text).strip()
        finally:
            document.Close(False)
    finally:
        if application is not None:
            application.Quit()
        path.unlink(missing_ok=True)


def _extract_legacy_excel_text(content: bytes, suffix: str) -> str:
    import tempfile

    import win32com.client

    with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as source:
        source.write(content)
        path = Path(source.name)
    application = None
    try:
        application = win32com.client.DispatchEx('Excel.Application')
        application.Visible = False
        application.DisplayAlerts = False
        application.AskToUpdateLinks = False
        workbook = application.Workbooks.Open(
            str(path.resolve()),
            UpdateLinks=0,
            ReadOnly=True,
            AddToMru=False,
        )
        try:
            blocks = []
            for sheet in workbook.Worksheets:
                used = sheet.UsedRange
                values = used.Value
                if values is None:
                    continue
                rows = values if isinstance(values, tuple) else ((values,),)
                if rows and not isinstance(rows[0], tuple):
                    rows = (rows,)
                blocks.append(f'工作表：{sheet.Name}')
                for row in rows:
                    cells = [str(value).strip() for value in row if value is not None and str(value).strip()]
                    if cells:
                        blocks.append(' | '.join(cells))
            return '\n'.join(blocks)
        finally:
            workbook.Close(False)
    finally:
        if application is not None:
            application.Quit()
        path.unlink(missing_ok=True)


def extract_order_source_text(filename: str, content: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_ORDER_EXTENSIONS:
        raise ValueError('支持的订单资料格式：PDF、Word（.doc/.docx）、Excel（.xls/.xlsx/.xlsm）和图片。')
    if suffix in {'.pdf', '.png', '.jpg', '.jpeg'}:
        text = extract_lc_text(filename, content)
    elif suffix == '.doc':
        try:
            text = _extract_legacy_word_text(content)
        except ImportError as exc:
            raise RuntimeError('读取旧版 .doc 文件需要在 Windows 上安装 Microsoft Word 和 pywin32。') from exc
    elif suffix == '.docx':
        document = Document(BytesIO(content))
        blocks = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            blocks.extend(
                ' | '.join(cell.text.strip() for cell in row.cells if cell.text.strip())
                for row in table.rows
            )
        text = '\n'.join(block for block in blocks if block.strip())
    elif suffix in {'.xls', '.xlsm'}:
        try:
            text = _extract_legacy_excel_text(content, suffix)
        except ImportError as exc:
            raise RuntimeError(f'读取旧版 {suffix} 文件需要在 Windows 上安装 Microsoft Excel 和 pywin32。') from exc
    else:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        try:
            blocks = []
            for sheet in workbook.worksheets:
                if sheet.sheet_state != 'visible':
                    continue
                blocks.append(f'工作表：{sheet.title}')
                for row in sheet.iter_rows(values_only=True):
                    cells = [str(value).strip() for value in row if value is not None and str(value).strip()]
                    if cells:
                        blocks.append(' | '.join(cells))
            text = '\n'.join(blocks)
        finally:
            workbook.close()
    if not text.strip():
        raise ValueError('未能从订单资料中提取文字，请上传清晰、可识别的原件。')
    return text


def recognize_order_text(text: str) -> dict[str, Any]:
    if not text.strip():
        raise ValueError('订单资料没有可识别的文字。')
    response = _call_chat_completion([
        {
            'role': 'system',
            'content': (
                '你是外贸订单资料录入助手。只从输入原文提取明确出现的信息，不得猜测、补充或换算出原文未提供的事实。'
                '缺失字段返回空字符串，数字字段缺失返回 null。日期统一为 YYYY-MM-DD；货币使用 ISO 4217 代码。'
                '仅返回一个 JSON 对象，不要 Markdown。字段格式：'
                '{"order_no":"","customer_name_cn":"","customer_name_en":"","customer_address":"","contact_person":"","email":"","phone":"",'
                '"product_name_cn":"","product_name_en":"","hs_code":"","cas_no":"","product_specification":"","product_description_en":"",'
                '"packaging_method":"","quantity":null,"unit":"","unit_price":null,"currency":"","gross_weight_per_unit":null,'
                '"net_weight_per_unit":null,"volume_per_unit":null,"loading_port":"","destination_port":"","incoterm":"",'
                '"payment_terms":"","shipment_date":"","latest_shipment_date":"","marks_and_numbers":"","special_requirements":""}。'
            ),
        },
        {'role': 'user', 'content': text[:50000]},
    ])
    match = re.search(r'\{.*\}', response, flags=re.DOTALL)
    if match is None:
        raise RuntimeError('订单信息识别结果不是有效 JSON。')
    try:
        result = json.loads(match.group(0))
    except json.JSONDecodeError as exc:
        raise RuntimeError('订单信息识别结果无法解析。') from exc
    if not isinstance(result, dict):
        raise RuntimeError('订单信息识别结果格式不正确。')

    string_fields = {
        'order_no', 'customer_name_cn', 'customer_name_en', 'customer_address',
        'contact_person', 'email', 'phone', 'product_name_cn', 'product_name_en',
        'hs_code', 'cas_no', 'product_specification', 'product_description_en',
        'packaging_method', 'unit', 'currency', 'loading_port', 'destination_port',
        'incoterm', 'payment_terms', 'shipment_date', 'latest_shipment_date',
        'marks_and_numbers', 'special_requirements',
    }
    cleaned: dict[str, Any] = {}
    for field in string_fields:
        value = result.get(field)
        cleaned[field] = str(value).strip() if value is not None else ''
    for field in ('quantity', 'unit_price', 'gross_weight_per_unit', 'net_weight_per_unit', 'volume_per_unit'):
        value = result.get(field)
        if value is None or value == '':
            cleaned[field] = None
            continue
        try:
            cleaned[field] = float(str(value).replace(',', ''))
        except (TypeError, ValueError):
            cleaned[field] = None
    return cleaned
