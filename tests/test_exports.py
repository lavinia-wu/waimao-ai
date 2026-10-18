from __future__ import annotations

import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from config.settings import settings
from database.db import get_connection
from database.init_db import create_tables
from modules.customers import create_customer, export_customers_excel, list_customers
from modules.products import export_products_excel

OWNER_A = 7
OWNER_B = 8

# 客户导出列号（与 CUSTOMER_EXPORT_HEADERS 一致）
PHONE_COLUMN = 7


def read_sheet(data: bytes) -> tuple[list[str], list[list]]:
    """读取工作簿首个工作表：返回（表头, 数据行）。"""
    workbook = load_workbook(BytesIO(data))
    sheet = workbook.active
    headers = [str(cell.value) for cell in sheet[1]]
    rows = [[cell.value for cell in row] for row in sheet.iter_rows(min_row=2)]
    workbook.close()
    return headers, rows


def cell_formats(data: bytes, row: int) -> list[str]:
    workbook = load_workbook(BytesIO(data))
    sheet = workbook.active
    formats = [
        sheet.cell(row, column).number_format
        for column in range(1, sheet.max_column + 1)
    ]
    workbook.close()
    return formats


def customer_payload(**overrides) -> dict:
    payload = {
        'customer_name_cn': '测试客户',
        'customer_name_en': 'Test Customer',
        'company_address': 'Lahore, Pakistan',
        'contact_person': 'Ali',
        'email': 'ali@example.com',
        'phone': '03001234567',
        'destination_port': 'Karachi',
        'payment_terms': 'L/C at sight',
        'incoterm': 'FOB',
        'remark': '备注内容',
        'country_code': 'PK',
    }
    payload.update(overrides)
    return payload


class CustomerExportTestCase(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.database_path = Path(self.temp_dir.name) / 'customer_export.db'
        self.settings_patch = patch.object(settings, 'db_path', self.database_path)
        self.settings_patch.start()
        create_tables()
        with get_connection() as conn:
            conn.execute(
                """
                INSERT INTO users (id, username, display_name, password_hash)
                VALUES (?, 'sales-a', '业务员甲', 'hash'), (?, 'sales-b', '业务员乙', 'hash')
                """,
                (OWNER_A, OWNER_B),
            )
            conn.commit()
        create_customer(customer_payload(customer_name_cn='甲方客户'), OWNER_A)
        create_customer(customer_payload(customer_name_cn='乙方客户'), OWNER_B)

    def tearDown(self) -> None:
        self.settings_patch.stop()
        self.temp_dir.cleanup()

    def test_sales_export_contains_only_own_customers(self) -> None:
        _, data = export_customers_excel('', OWNER_A)
        _, rows = read_sheet(data)
        self.assertEqual(1, len(rows))
        self.assertEqual('甲方客户', rows[0][0])

    def test_admin_export_contains_all_customers_with_owner(self) -> None:
        file_name, data = export_customers_excel('', None)
        headers, rows = read_sheet(data)
        self.assertEqual(2, len(rows))
        self.assertIn('归属业务员', headers)
        owner_index = headers.index('归属业务员')
        self.assertEqual({'业务员甲', '业务员乙'}, {row[owner_index] for row in rows})
        self.assertTrue(file_name.startswith('客户_'))
        self.assertTrue(file_name.endswith('.xlsx'))

    def test_sales_export_has_no_owner_column(self) -> None:
        _, data = export_customers_excel('', OWNER_A)
        headers, _ = read_sheet(data)
        self.assertNotIn('归属业务员', headers)

    def test_filtered_export_follows_search(self) -> None:
        _, data = export_customers_excel('甲方', None)
        _, rows = read_sheet(data)
        self.assertEqual(1, len(rows))
        self.assertEqual('甲方客户', rows[0][0])

    def test_phone_kept_as_text(self) -> None:
        create_customer(customer_payload(customer_name_cn='前导零客户', phone='0012345'), OWNER_A)
        _, data = export_customers_excel('前导零', OWNER_A)
        _, rows = read_sheet(data)
        self.assertEqual('0012345', rows[0][PHONE_COLUMN - 1])
        self.assertEqual('@', cell_formats(data, 2)[PHONE_COLUMN - 1])

    def test_formula_injection_is_escaped(self) -> None:
        create_customer(customer_payload(customer_name_cn='=1+1'), OWNER_A)
        _, data = export_customers_excel('=1+1', OWNER_A)
        _, rows = read_sheet(data)
        self.assertEqual("'=1+1", rows[0][0])
        self.assertNotEqual('=1+1', rows[0][0])

    def test_no_sensitive_or_internal_columns(self) -> None:
        _, data = export_customers_excel('', None)
        headers, rows = read_sheet(data)
        forbidden = ('id', 'user_id', 'password_hash', 'updated_by', 'path')
        for column in forbidden:
            self.assertNotIn(column, headers)
        for row in rows:
            for value in row:
                self.assertNotIn('C:\\', str(value))
                self.assertNotIn('/uploads/', str(value))

    def test_export_does_not_modify_data(self) -> None:
        before = list_customers('', None)
        export_customers_excel('', None)
        export_customers_excel('', OWNER_A)
        after = list_customers('', None)
        self.assertEqual(before, after)


class ProductExportTestCase(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.database_path = Path(self.temp_dir.name) / 'product_export.db'
        self.settings_patch = patch.object(settings, 'db_path', self.database_path)
        self.settings_patch.start()
        create_tables()
        with get_connection() as conn:
            conn.execute(
                """
                INSERT INTO products (
                    product_name_en, product_name_cn, sku, hs_code, cas_no,
                    unit_price, currency, packaging_method, specification,
                    stock_qty, unit, stock_updated_at, msds_file_path
                ) VALUES
                    ('PP 550J', '聚丙烯', '0012', '3902.10', '0000111',
                     1200.5, 'USD', '25kg bag', '注塑级', 12.0, 'MT',
                     '2026-10-01 10:00:00', 'C:/secret/msds/pp550j.pdf'),
                    ('HDPE 6098H', '高密度聚乙烯', '0013', '3901.20', '0000222',
                     NULL, 'USD', '25kg bag', '吹塑级', NULL, 'MT', NULL, NULL),
                    ('LDPE 2426H', '低密度聚乙烯', '0014', '3901.10', '0000333',
                     900, 'USD', '25kg bag', '薄膜级', 0, 'MT',
                     '2026-10-02 10:00:00', NULL)
                """
            )
            conn.commit()

    def tearDown(self) -> None:
        self.settings_patch.stop()
        self.temp_dir.cleanup()

    def test_codes_kept_as_text(self) -> None:
        _, data = export_products_excel('')
        headers, rows = read_sheet(data)
        formats = cell_formats(data, 2)
        for column_name in ('产品编码(SKU)', 'HS编码', 'CAS号'):
            index = headers.index(column_name)
            self.assertEqual('@', formats[index], f'{column_name} 未按文本保存')
        sku_index = headers.index('产品编码(SKU)')
        # 行序由 list_products 决定（updated_at DESC），断言不依赖行序
        self.assertEqual(
            ['0012', '0013', '0014'], sorted(str(row[sku_index]) for row in rows)
        )

    def test_stock_unmaintained_differs_from_zero(self) -> None:
        _, data = export_products_excel('')
        headers, rows = read_sheet(data)
        stock_index = headers.index('库存数量')
        values = {row[headers.index('中文品名')]: row[stock_index] for row in rows}
        self.assertEqual(12.0, values['聚丙烯'])
        self.assertEqual('未维护', values['高密度聚乙烯'])
        self.assertEqual(0, values['低密度聚乙烯'])
        self.assertNotEqual(values['高密度聚乙烯'], values['低密度聚乙烯'])

    def test_no_file_path_or_binary_content(self) -> None:
        file_name, data = export_products_excel('')
        headers, rows = read_sheet(data)
        msds_index = headers.index('MSDS资料')
        name_index = headers.index('中文品名')
        by_name = {row[name_index]: row[msds_index] for row in rows}
        self.assertEqual('有', by_name['聚丙烯'])
        self.assertEqual('无', by_name['高密度聚乙烯'])
        self.assertEqual('无', by_name['低密度聚乙烯'])
        for row in rows:
            for value in row:
                self.assertNotIn('C:/secret/msds', str(value))
        self.assertIn('产品_', file_name)

    def test_export_does_not_modify_stock(self) -> None:
        _, data = export_products_excel('')
        with get_connection() as conn:
            before = conn.execute(
                'SELECT sku, stock_qty, stock_updated_at FROM products ORDER BY id'
            ).fetchall()
        export_products_excel('')
        with get_connection() as conn:
            after = conn.execute(
                'SELECT sku, stock_qty, stock_updated_at FROM products ORDER BY id'
            ).fetchall()
        self.assertEqual([tuple(row) for row in before], [tuple(row) for row in after])
        self.assertGreater(len(data), 0)


if __name__ == '__main__':
    unittest.main()
