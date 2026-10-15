from __future__ import annotations

from pathlib import Path
import shutil
from typing import Any, Dict, Iterable

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from services.amount_words import amount_in_english_words
from config.settings import settings


DEFAULT_TEMPLATE_FILES = {
    'PI': 'PI_Professional_template.docx',
    'Commercial Invoice': 'Commercial_Invoice_Professional_template.docx',
    'Packing List': 'Packing_List_Professional_template.docx',
    'Sales Contract': 'Sales_Contract_Professional_template.docx',
}

LEGACY_BUILTIN_TEMPLATE_FILES = {
    'PI_template.docx',
    'Commercial_Invoice_template.docx',
    'Packing_List_template.docx',
    'Sales_Contract_template.docx',
}


def get_default_template_path(doc_type: str) -> Path:
    try:
        filename = DEFAULT_TEMPLATE_FILES[doc_type]
    except KeyError as exc:
        raise ValueError(f'不支持的单证类型：{doc_type}') from exc
    return settings.system_templates_dir / filename


def list_word_templates() -> list[Path]:
    settings.templates_dir.mkdir(parents=True, exist_ok=True)
    return sorted(
        (
            path for path in settings.templates_dir.rglob('*.docx')
            if path.is_file()
            and path.name not in LEGACY_BUILTIN_TEMPLATE_FILES
            and path.name not in DEFAULT_TEMPLATE_FILES.values()
        ),
        key=lambda path: path.name.lower(),
    )


def _replace_placeholder_in_runs(paragraph, old_text: str, new_value: str) -> None:
    runs = paragraph.runs
    search_offset = 0
    while True:
        combined_text = ''.join(run.text for run in runs)
        match_start = combined_text.find(old_text, search_offset)
        if match_start < 0:
            break
        match_end = match_start + len(old_text)
        cumulative = 0
        start_index = end_index = -1
        start_offset = end_offset = 0
        for index, run in enumerate(runs):
            next_cumulative = cumulative + len(run.text)
            if start_index == -1 and match_start < next_cumulative:
                start_index = index
                start_offset = match_start - cumulative
            if match_end <= next_cumulative:
                end_index = index
                end_offset = match_end - cumulative
                break
            cumulative = next_cumulative

        if start_index < 0 or end_index < 0:
            break
        if start_index == end_index:
            run = runs[start_index]
            run.text = run.text[:start_offset] + new_value + run.text[end_offset:]
        else:
            runs[start_index].text = runs[start_index].text[:start_offset] + new_value
            for index in range(start_index + 1, end_index):
                runs[index].text = ''
            runs[end_index].text = runs[end_index].text[end_offset:]
        search_offset = match_start + len(new_value)


def _replace_placeholders_in_element(element, values: Dict[str, Any]) -> None:
    for paragraph in element.paragraphs:
        for key, value in values.items():
            placeholder = '{{' + key + '}}'
            if placeholder in paragraph.text:
                _replace_placeholder_in_runs(paragraph, placeholder, str(value))
    for table in element.tables:
        for row in table.rows:
            for cell in row.cells:
                _replace_placeholders_in_element(cell, values)


def _element_contains_placeholder(element, placeholder: str) -> bool:
    for paragraph in element.paragraphs:
        if placeholder in paragraph.text:
            return True
    for table in element.tables:
        for row in table.rows:
            for cell in row.cells:
                if _element_contains_placeholder(cell, placeholder):
                    return True
    return False


def _shade_cell(cell, fill: str) -> None:
    shading = cell._tc.get_or_add_tcPr()
    element = shading.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}shd')
    if element is None:
        from docx.oxml import OxmlElement

        element = OxmlElement('w:shd')
        shading.append(element)
    element.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fill', fill)


def _format_table(table, header: bool = True) -> None:
    table.style = 'Table Grid'
    table.autofit = True
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(3)
                paragraph.paragraph_format.space_before = Pt(3)
                for run in paragraph.runs:
                    run.font.name = 'Arial'
                    run.font.size = Pt(9)
            if header and row_index == 0:
                _shade_cell(cell, '17365D')
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)


def _add_table(document: Document, headers: list[str], rows: list[list[str]], header: bool = True):
    table = document.add_table(rows=1 if header else 0, cols=len(headers))
    if header:
        for cell, text in zip(table.rows[0].cells, headers):
            cell.text = text
    for row_values in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, row_values):
            cell.text = text
    _format_table(table, header=header)
    return table


def _add_title(document: Document, title: str, subtitle: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run('{{company_name}}')
    run.bold = True
    run.font.name = 'Arial'
    run.font.size = Pt(15)
    run.font.color.rgb = RGBColor(31, 78, 121)

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(12)
    run = paragraph.add_run(title)
    run.bold = True
    run.font.name = 'Arial'
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(31, 78, 121)
    paragraph.add_run(f'\n{subtitle}').font.size = Pt(9)


def _build_builtin_template(doc_type: str, output_path: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)
    normal = document.styles['Normal']
    normal.font.name = 'Arial'
    normal.font.size = Pt(9)
    normal.paragraph_format.space_after = Pt(4)

    if doc_type == 'PI':
        _add_title(document, 'PROFORMA INVOICE', 'COMMERCIAL OFFER')
        _add_table(document, ['SELLER / BENEFICIARY', 'BUYER / APPLICANT'], [[
            '{{company_name}}\n{{company_address}}\n{{company_email}}\n{{company_phone}}',
            '{{customer_name_en}}\n{{customer_address}}\nAttn: {{contact_person}}\n{{email}}',
        ]])
        document.add_paragraph()
        _add_table(document, ['PI NO.', 'DATE', 'ORDER NO.'], [['{{document_no}}', '{{date}}', '{{order_no}}']])
        document.add_heading('Product and Price', level=2)
        _add_table(document, ['NO.', 'DESCRIPTION OF GOODS', 'QUANTITY', 'UNIT PRICE', 'AMOUNT'], [[
            '1', '{{product_name_en}}', '{{quantity}} {{unit}}', '{{currency}} {{unit_price}}', '{{currency}} {{total_amount}}',
        ]])
        document.add_paragraph('AMOUNT IN WORDS: {{amount_in_words}}')
        document.add_heading('Commercial Terms', level=2)
        _add_table(document, ['DELIVERY TERM', 'LOADING PORT', 'DESTINATION PORT', 'PAYMENT TERMS'], [[
            '{{incoterm}}', '{{loading_port}}', '{{destination_port}}', '{{payment_terms}}',
        ]])
        document.add_paragraph('Shipment date: {{shipment_date}}    Latest shipment date: {{latest_shipment_date}}')
        document.add_paragraph('Remarks: {{remarks}}')
        document.add_paragraph('L/C requirements: {{lc_document_instruction}}')
        document.add_paragraph('\nFor and on behalf of the Seller')
        document.add_paragraph('Authorized Signature: __________________________    Date: ______________')
    elif doc_type == 'Commercial Invoice':
        _add_title(document, 'COMMERCIAL INVOICE', 'EXPORT INVOICE')
        _add_table(document, ['SELLER / EXPORTER', 'BUYER / IMPORTER'], [[
            '{{company_name}}\n{{company_address}}\n{{company_email}}\n{{company_phone}}',
            '{{customer_name_en}}\n{{customer_address}}\nAttn: {{contact_person}}\n{{email}}',
        ]])
        document.add_paragraph()
        _add_table(document, ['INVOICE NO.', 'INVOICE DATE', 'ORDER / CONTRACT NO.', 'CURRENCY'], [[
            '{{document_no}}', '{{date}}', '{{order_no}}', '{{currency}}',
        ]])
        _add_table(document, ['LOADING PORT', 'PORT OF DESTINATION', 'INCOTERM', 'PAYMENT TERMS'], [[
            '{{loading_port}}', '{{destination_port}}', '{{incoterm}}', '{{payment_terms}}',
        ]])
        document.add_heading('Description of Goods', level=2)
        _add_table(document, ['NO.', 'DESCRIPTION', 'QUANTITY', 'UNIT PRICE', 'LINE TOTAL'], [[
            '1', '{{product_name_en}}', '{{quantity}} {{unit}}', '{{unit_price}}', '{{total_amount}}',
        ]])
        _add_table(document, ['TOTAL IN {{currency}}', '{{total_amount}}'], [], header=False)
        document.add_paragraph('AMOUNT IN WORDS: {{amount_in_words}}')
        document.add_paragraph('Marks & Numbers: {{marks_and_numbers}}')
        document.add_paragraph('Gross weight: {{gross_weight}} kg    Net weight: {{net_weight}} kg    Volume: {{volume}} CBM')
        document.add_paragraph('L/C requirements: {{lc_document_instruction}}')
        document.add_paragraph('\nWe certify that the information stated above is true and correct.')
        document.add_paragraph('Authorized Signature: __________________________    Company Stamp')
    elif doc_type == 'Packing List':
        _add_title(document, 'PACKING LIST', 'SHIPPING AND PACKING DETAILS')
        _add_table(document, ['SHIPPER / EXPORTER', 'CONSIGNEE / BUYER'], [[
            '{{company_name}}\n{{company_address}}\n{{company_email}}',
            '{{customer_name_en}}\n{{customer_address}}\nAttn: {{contact_person}}',
        ]])
        document.add_paragraph()
        _add_table(document, ['PACKING LIST NO.', 'DATE', 'INVOICE / ORDER NO.'], [[
            '{{document_no}}', '{{date}}', '{{order_no}}',
        ]])
        _add_table(document, ['PORT OF LOADING', 'PORT OF DESTINATION', 'MARKS & NUMBERS'], [[
            '{{loading_port}}', '{{destination_port}}', '{{marks_and_numbers}}',
        ]])
        document.add_heading('Packing Details', level=2)
        _add_table(document, ['NO.', 'GOODS DESCRIPTION', 'PACKAGES / QUANTITY', 'NET WEIGHT (KG)', 'GROSS WEIGHT (KG)', 'MEASUREMENT (CBM)'], [[
            '1', '{{product_name_en}}', '{{quantity}} {{unit}}', '{{net_weight}}', '{{gross_weight}}', '{{volume}}',
        ]])
        _add_table(document, ['TOTAL PACKAGES / QUANTITY', 'TOTAL NET WEIGHT', 'TOTAL GROSS WEIGHT', 'TOTAL MEASUREMENT'], [[
            '{{quantity}} {{unit}}', '{{net_weight}} kg', '{{gross_weight}} kg', '{{volume}} CBM',
        ]])
        document.add_paragraph('Remarks: {{remarks}}')
        document.add_paragraph('L/C requirements: {{lc_document_instruction}}')
        document.add_paragraph('\nPacking details are true and correct.')
        document.add_paragraph('Authorized Signature: __________________________    Company Stamp')
    else:
        _add_title(document, 'SALES CONTRACT', 'INTERNATIONAL SALES AGREEMENT')
        document.add_paragraph(
            'Contract No.: {{document_no}}        Date: {{date}}\n'
            'Seller: {{company_name}}, {{company_address}}\n'
            'Buyer: {{customer_name_en}}, {{customer_address}}'
        )
        document.add_paragraph('The Seller and the Buyer agree to the following terms and conditions:')
        clauses = [
            ('1. GOODS AND QUANTITY', 'Product: {{product_name_en}}    Quantity: {{quantity}} {{unit}}'),
            ('2. PRICE AND CONTRACT VALUE', 'Unit price: {{currency}} {{unit_price}}    Total contract value: {{currency}} {{total_amount}}'),
            ('3. DELIVERY', 'Trade term: {{incoterm}}    Loading port: {{loading_port}}    Destination port: {{destination_port}}\n'
             'Shipment date: {{shipment_date}}    Latest shipment date: {{latest_shipment_date}}'),
            ('4. PAYMENT', '{{payment_terms}}'),
            ('5. PACKING AND SHIPPING MARKS', '{{marks_and_numbers}}\nGross weight: {{gross_weight}} kg    Net weight: {{net_weight}} kg    Volume: {{volume}} CBM'),
            ('6. SPECIAL TERMS', '{{remarks}}\n{{lc_document_instruction}}'),
        ]
        for heading, body in clauses:
            document.add_heading(heading, level=2)
            document.add_paragraph(body)
        document.add_paragraph('\nSELLER                                      BUYER')
        document.add_paragraph('Authorized Signature: __________________    Authorized Signature: __________________')
        document.add_paragraph('Name / Title: __________________________    Name / Title: __________________________')
        document.add_paragraph('Date: __________________________________    Date: __________________________________')

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('{{company_name}}  |  {{document_no}}').font.size = Pt(8)
    document.save(output_path)


def _set_paragraph_text(paragraph, text: str) -> None:
    runs = paragraph.runs
    if runs:
        runs[0].text = ''
        for run in runs[1:]:
            run.text = ''
        target_run = runs[0]
    else:
        target_run = paragraph.add_run()
    for index, line in enumerate(str(text).split('\n')):
        if index:
            target_run.add_break()
        target_run.add_text(line)


def _set_cell_text(cell, text: str) -> None:
    paragraphs = cell.paragraphs
    if not paragraphs:
        cell.add_paragraph()
        paragraphs = cell.paragraphs
    _set_paragraph_text(paragraphs[0], text)
    paragraph_format = paragraphs[0].paragraph_format
    paragraph_format.space_before = Pt(0)
    paragraph_format.space_after = Pt(0)
    paragraph_format.line_spacing = 1
    for paragraph in paragraphs[1:]:
        cell._tc.remove(paragraph._p)


def _currency_symbol(currency: str) -> str | None:
    return {
        'USD': '$',
        'EUR': '€',
        'GBP': '£',
        'CNY': '¥',
        'JPY': '¥',
        'AUD': 'A$',
        'CAD': 'C$',
        'HKD': 'HK$',
        'SGD': 'S$',
    }.get(currency.strip().upper())


def _fill_uploaded_proforma(document: Document, values: Dict[str, Any]) -> None:
    if len(document.paragraphs) > 6:
        address_lines = str(values.get('company_address', '')).splitlines()
        paragraph_values = {
            0: values.get('company_name', ''),
            1: address_lines[0] if address_lines else '',
            2: address_lines[1] if len(address_lines) > 1 else '',
            3: address_lines[2] if len(address_lines) > 2 else '',
            6: values.get('date', ''),
        }
        for index, value in paragraph_values.items():
            _set_paragraph_text(document.paragraphs[index], str(value or ''))

    if not document.tables:
        return
    table = document.tables[0]
    if len(table.rows) < 9:
        return
    buyer = '\n'.join(filter(None, [
        str(values.get('customer_name_en', '')).strip(),
        str(values.get('customer_address', '')).strip(),
        str(values.get('email', '')).strip(),
    ]))
    _set_cell_text(table.cell(0, 0), f'SOLD TO:\n{buyer}')
    _set_cell_text(table.cell(0, 3), f"Contract No. {values.get('order_no', '')}")
    _set_cell_text(table.cell(1, 0), f'CONSIGNEE:\n{buyer}')
    _set_cell_text(table.cell(2, 0), f"PROFORMA INVOICE DATE: {values.get('date', '')}")
    _set_cell_text(table.cell(3, 0), f"Origin: {values.get('loading_port', '')}")
    _set_cell_text(table.cell(3, 2), f"Delivery: {values.get('shipment_date', '')}")

    gross_kg = float(values.get('gross_weight') or 0)
    net_kg = float(values.get('net_weight') or 0)
    gross_mt = gross_kg / 1000
    net_mt = net_kg / 1000
    currency = str(values.get('currency', ''))
    currency_symbol = _currency_symbol(currency)
    product_code = str(values.get('product_code') or '').strip()
    hs_code = str(values.get('hs_code') or '').strip()
    if product_code and hs_code and product_code == hs_code:
        product_code = ''
    description = '\n'.join(filter(None, [
        str(values.get('product_name_en', '')).strip(),
        str(values.get('product_specification', '')).strip(),
        f"PACKING: {values.get('packaging_method')}" if values.get('packaging_method') else '',
        f"HS CODE: {values.get('hs_code')}" if values.get('hs_code') else '',
    ]))
    for cell_index, value in {
        0: product_code,
        1: description,
        3: f'{gross_kg:g}',
        4: f'{gross_mt:g}',
        5: currency_symbol or '',
        6: values.get('unit_price', '') if currency_symbol else f"{currency} {values.get('unit_price', '')}".strip(),
        7: currency_symbol or '',
        8: values.get('total_amount', '') if currency_symbol else f"{currency} {values.get('total_amount', '')}".strip(),
    }.items():
        _set_cell_text(table.cell(5, cell_index), str(value or ''))
    for cell_index, value in {
        3: f'{gross_kg:g}',
        4: f'{gross_mt:g}',
        5: values.get('incoterm', ''),
        7: currency_symbol or '',
        8: values.get('total_amount', '') if currency_symbol else f"{currency} {values.get('total_amount', '')}".strip(),
    }.items():
        _set_cell_text(table.cell(6, cell_index), str(value or ''))
    _set_cell_text(
        table.cell(7, 0),
        f"AMOUNT IN WORDS: {amount_in_english_words(values.get('total_amount'), currency)}",
    )
    additions = []
    lc_instruction = str(values.get('lc_document_instruction', '')).strip()
    remarks = str(values.get('remarks', '')).strip()
    if lc_instruction:
        additions.append(('Letter of Credit Requirements:', lc_instruction))
    if remarks:
        additions.append(('Order Remarks:', remarks))
    if additions:
        terms_cell = table.cell(8, 0)
        for heading, text in additions:
            paragraph = terms_cell.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.add_run(heading)
            for line in text.splitlines():
                paragraph = terms_cell.add_paragraph(line)
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)


def _fill_uploaded_sales_contract(document: Document, values: Dict[str, Any]) -> None:
    paragraphs = document.paragraphs
    if len(paragraphs) > 3:
        introduction = paragraphs[3]
        _replace_between(introduction, 'between ', ', with office address at', str(values.get('company_name', '')))
        _replace_between(
            introduction,
            'with office address at',
            ' (hereinafter referred to as the "SELLER")',
            f" {values.get('company_address', '')}",
        )
        seller_marker = '(hereinafter referred to as the "SELLER") and '
        buyer_start = introduction.text.find(seller_marker)
        if buyer_start >= 0:
            buyer_start += len(seller_marker)
            buyer_end = introduction.text.find('addressed at', buyer_start)
            if buyer_end > buyer_start:
                _replace_paragraph_text(
                    introduction,
                    introduction.text[buyer_start:buyer_end],
                    str(values.get('customer_name_en', '')),
                )
        if 'addressed at' in introduction.text:
            _replace_paragraph_text(introduction, 'addressed at', ' addressed at')
        _replace_between(
            introduction,
            'addressed at',
            ' (hereinafter referred to as the "BUYER")',
            f" {values.get('customer_address', '')}",
        )
        _replace_between(
            introduction,
            'pay for the ',
            ' enumerated',
            str(values.get('product_name_en', '')),
        )

    net_weight_mt = float(values.get('net_weight') or 0) / 1000
    fields = [
        (('1.', 'CONTRACT'), str(values.get('document_no', ''))),
        (('2.', 'PRODUCT'), str(values.get('product_name_en', ''))),
        (('GRADE',), str(values.get('product_specification', ''))),
        (('4.', 'NET QUANTITY'), f'{net_weight_mt:g}MTS'),
        (('5.', 'SPECIFICATIONS'), str(values.get('product_specification', ''))),
        (
            ('6.', 'PRICE / VAT Inclusive'),
            f"{values.get('currency', '')} {values.get('unit_price', '')}".strip(),
        ),
        (('7.', 'PACKAGING'), str(values.get('packaging_method', ''))),
        (('8.', 'PAYMENT'), str(values.get('payment_terms', ''))),
        (('9.', 'DELIVERY'), str(values.get('shipment_date', ''))),
        (('LOADING PORT',), str(values.get('loading_port', ''))),
        (('DESTINATION',), str(values.get('destination_port', ''))),
    ]
    for prefixes, value in fields:
        paragraph = next(
            (
                item for item in paragraphs
                if item.text.strip().upper().startswith(prefixes[0])
                and all(prefix in item.text.upper() for prefix in prefixes[1:])
            ),
            None,
        )
        if paragraph is not None:
            _set_value_after_label(paragraph, prefixes[-1], value)

    for index, value in ((33, values.get('company_name')), (41, values.get('customer_name_en')), (48, values.get('company_name'))):
        if index < len(paragraphs) and value:
            _replace_paragraph_text(paragraphs[index], paragraphs[index].text, str(value))

    if len(paragraphs) > 51 and values.get('company_address'):
        address_lines = str(values['company_address']).splitlines()
        for index in range(49, 52):
            line = address_lines[index - 49] if index - 49 < len(address_lines) else ''
            _replace_paragraph_text(paragraphs[index], paragraphs[index].text, line)


def _fill_sample_populated_template(document: Document, doc_type: str, values: Dict[str, Any], source: Path) -> None:
    filename = source.name.lower()
    if doc_type == 'PI' and ('proforma' in filename or 'proforma invoice' in filename):
        _fill_uploaded_proforma(document, values)
    elif (
        doc_type == 'Sales Contract'
        and 'sales contract' in filename
        and not _element_contains_placeholder(document, '{{')
    ):
        _fill_uploaded_sales_contract(document, values)


def _normalize_word_fonts(document: Document, uniform_contract: bool = False) -> None:
    from docx.oxml.ns import qn

    def normalize_run_properties(properties) -> None:
        for tag in (
            'w:b',
            'w:bCs',
            'w:i',
            'w:iCs',
            'w:strike',
            'w:dstrike',
            'w:caps',
            'w:smallCaps',
            'w:outline',
            'w:shadow',
            'w:emboss',
            'w:imprint',
        ):
            for element in list(properties.findall(qn(tag))):
                properties.remove(element)
        for tag in ('w:spacing', 'w:kern', 'w:position', 'w:w'):
            for element in list(properties.findall(qn(tag))):
                properties.remove(element)
        underline = properties.find(qn('w:u'))
        if underline is not None:
            properties.remove(underline)

    def set_font_attributes(element) -> None:
        for fonts in element.iter(qn('w:rFonts')):
            fonts.set(qn('w:ascii'), 'Arial')
            fonts.set(qn('w:hAnsi'), 'Arial')
            fonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
            fonts.set(qn('w:cs'), 'Arial')
            for theme_attribute in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
                fonts.attrib.pop(qn(f'w:{theme_attribute}'), None)

    def normalize_paragraph(paragraph) -> None:
        for run in paragraph.runs:
            run.font.name = 'Arial'
            if uniform_contract:
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0, 0, 0)
                run.bold = False
                run.italic = False
                run.underline = False
                run.font.strike = False
                normalize_run_properties(run._r.get_or_add_rPr())
        if uniform_contract:
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.line_spacing = 1

    def normalize_element(element) -> None:
        for paragraph in element.paragraphs:
            normalize_paragraph(paragraph)
        for table in element.tables:
            for row in table.rows:
                for cell in row.cells:
                    normalize_element(cell)

    for style in document.styles:
        if hasattr(style, 'font'):
            style.font.name = 'Arial'
            if uniform_contract:
                style.font.size = Pt(9)
                style.font.color.rgb = RGBColor(0, 0, 0)
                style.font.bold = False
                style.font.italic = False
                style.font.underline = False
                style.font.strike = False

    normalize_element(document)
    for section in document.sections:
        normalize_element(section.header)
        normalize_element(section.footer)
    if uniform_contract:
        for element in (
            document._element,
            document.styles.element,
            document.part.numbering_part.element,
        ):
            for properties in element.iter(qn('w:rPr')):
                normalize_run_properties(properties)
        numbering = document.part.numbering_part.element
        for properties in numbering.iter(qn('w:rPr')):
            normalize_run_properties(properties)
            fonts = properties.find(qn('w:rFonts'))
            if fonts is None:
                from docx.oxml import OxmlElement

                fonts = OxmlElement('w:rFonts')
                properties.insert(0, fonts)
            fonts.set(qn('w:ascii'), 'Arial')
            fonts.set(qn('w:hAnsi'), 'Arial')
            fonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
            fonts.set(qn('w:cs'), 'Arial')
            size = properties.find(qn('w:sz'))
            if size is None:
                from docx.oxml import OxmlElement

                size = OxmlElement('w:sz')
                properties.append(size)
            size.set(qn('w:val'), '18')
            color = properties.find(qn('w:color'))
            if color is None:
                from docx.oxml import OxmlElement

                color = OxmlElement('w:color')
                properties.append(color)
            color.set(qn('w:val'), '000000')
    set_font_attributes(document._element)
    set_font_attributes(document.styles.element)
    for section in document.sections:
        set_font_attributes(section.header._element)
        set_font_attributes(section.footer._element)


def _replace_paragraph_text(paragraph, old_text: str, new_text: str) -> None:
    if old_text and old_text in paragraph.text:
        _replace_placeholder_in_runs(paragraph, old_text, new_text)


def _replace_between(
    paragraph,
    start_marker: str,
    end_marker: str,
    new_value: str,
) -> None:
    text = paragraph.text
    start = text.find(start_marker)
    if start < 0:
        return
    start += len(start_marker)
    end = text.find(end_marker, start)
    if end > start:
        _replace_paragraph_text(paragraph, text[start:end], new_value)


def _set_value_after_label(paragraph, label: str, value: str) -> None:
    text = paragraph.text
    label_end = text.lower().find(label.lower())
    if label_end < 0:
        return
    label_end += len(label)
    tail = text[label_end:]
    leading = tail[:len(tail) - len(tail.lstrip(' \t\u00a0'))]
    trailing = tail[len(tail.rstrip(' \t\u00a0')):]
    old_tail = tail
    if old_tail:
        _replace_paragraph_text(paragraph, old_tail, f'{leading}{value}{trailing}')


def _normalize_uploaded_proforma_layout(document: Document) -> None:
    from docx.oxml.ns import qn
    from docx.shared import Inches

    body = document._element.body
    for paragraph in list(body.findall(qn('w:p'))):
        properties = paragraph.find(qn('w:pPr'))
        if properties is None:
            continue
        section_properties = properties.find(qn('w:sectPr'))
        if section_properties is None:
            continue
        properties.remove(section_properties)
        if not paragraph.xpath('.//w:t'):
            body.remove(paragraph)

    children = list(body)
    table_positions = [index for index, child in enumerate(children) if child.tag == qn('w:tbl')]
    if table_positions:
        for child in children[table_positions[-1] + 1:]:
            if child.tag == qn('w:p') and not child.xpath('.//w:t'):
                body.remove(child)

    section = document.sections[-1]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.38)
    section.right_margin = Inches(0.38)
    section.header_distance = Inches(0.2)
    section.footer_distance = Inches(0.2)


def ensure_default_template(doc_type: str, template_path: str | None = None) -> str:
    output_path = Path(template_path) if template_path else get_default_template_path(doc_type)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    if output_path.exists():
        return str(output_path)

    legacy_path = settings.templates_dir / output_path.name
    if template_path is None and legacy_path.is_file():
        shutil.move(str(legacy_path), str(output_path))
        return str(output_path)

    _build_builtin_template(doc_type, output_path)
    return str(output_path)


def export_document_word(doc_type: str, values: Dict[str, Any], template_path: str | None = None, output_path: str | None = None) -> str:
    selected_template = template_path or ensure_default_template(doc_type)
    source = Path(selected_template)
    if not source.exists():
        raise FileNotFoundError(f'未找到所选 Word 模板：{source}')
    if source.suffix.lower() != '.docx':
        raise ValueError('Word 单证模板必须是 .docx 文件。')

    word_values = dict(values)
    if doc_type in {'PI', 'Commercial Invoice'}:
        word_values['amount_in_words'] = amount_in_english_words(
            values.get('total_amount'),
            str(values.get('currency') or 'USD'),
        )
    document = Document(source)
    _fill_sample_populated_template(document, doc_type, word_values, source)
    has_amount_words_placeholder = _element_contains_placeholder(document, '{{amount_in_words}}')
    has_lc_placeholder = _element_contains_placeholder(document, '{{lc_document_instruction}}')
    for section in document.sections:
        has_lc_placeholder = has_lc_placeholder or _element_contains_placeholder(section.header, '{{lc_document_instruction}}')
        has_lc_placeholder = has_lc_placeholder or _element_contains_placeholder(section.footer, '{{lc_document_instruction}}')
    _replace_placeholders_in_element(document, word_values)
    for section in document.sections:
        _replace_placeholders_in_element(section.header, word_values)
        _replace_placeholders_in_element(section.footer, word_values)
    if (
        doc_type in {'PI', 'Commercial Invoice'}
        and source.name in DEFAULT_TEMPLATE_FILES.values()
        and not has_amount_words_placeholder
    ):
        document.add_paragraph(f"AMOUNT IN WORDS: {word_values['amount_in_words']}")
    if doc_type == 'PI' and 'proforma' in source.name.lower():
        _normalize_uploaded_proforma_layout(document)
    lc_instruction = str(word_values.get('lc_document_instruction', '')).strip()
    if lc_instruction and not has_lc_placeholder:
        heading = document.add_paragraph()
        heading_run = heading.add_run('Letter of Credit Requirements')
        heading_run.bold = True
        heading_run.font.size = Pt(12)
        for line in lc_instruction.splitlines():
            document.add_paragraph(line)
    if doc_type != 'Sales Contract':
        _normalize_word_fonts(document)

    target = Path(output_path) if output_path else settings.outputs_dir / values.get('order_no', 'document') / f'{doc_type}_export.docx'
    target.parent.mkdir(parents=True, exist_ok=True)
    document.save(target)
    return str(target)
