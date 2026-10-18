from __future__ import annotations

from copy import copy
from datetime import datetime
from io import BytesIO
import os
from pathlib import Path
import tempfile
from typing import Any, Dict

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

from services.amount_words import amount_in_english_words
from config.settings import settings


NAVY = '17365D'
BLUE = '1F4E79'
PALE_BLUE = 'EAF1F8'
WHITE = 'FFFFFF'
GRID = 'B7C9D6'


def list_excel_templates() -> list[Path]:
    settings.templates_dir.mkdir(parents=True, exist_ok=True)
    return sorted(
        (
            path for path in settings.templates_dir.rglob('*.xlsx')
            if path.is_file() and not path.name.startswith('~$')
        ),
        key=lambda path: path.name.lower(),
    )


def get_excel_template_sheets(template_path: str) -> list[str]:
    source = Path(template_path)
    if not source.exists():
        raise FileNotFoundError(f'未找到所选 Excel 模板：{source}')
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        return [sheet.title for sheet in workbook.worksheets if sheet.sheet_state == 'visible']
    finally:
        workbook.close()


def _set_template_cell(sheet, address: str, value: Any) -> None:
    if value is None or value == '':
        return
    cell = sheet[address]
    if cell.__class__.__name__ == 'MergedCell':
        raise ValueError(f'模板字段位置 {sheet.title}!{address} 是合并单元格中的非起始格。')
    cell.value = value
    cell.alignment = Alignment(
        horizontal=cell.alignment.horizontal,
        vertical=cell.alignment.vertical or 'center',
        wrap_text=True,
    )


def _normalize_workbook_fonts(workbook, font_name: str = 'Arial') -> None:
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.__class__.__name__ == 'MergedCell':
                    continue
                font = copy(cell.font)
                font.name = font_name
                cell.font = font
                if isinstance(cell.value, str) and '\n' in cell.value:
                    cell.alignment = Alignment(
                        horizontal=cell.alignment.horizontal,
                        vertical=cell.alignment.vertical or 'center',
                        text_rotation=cell.alignment.text_rotation,
                        wrap_text=True,
                        shrink_to_fit=cell.alignment.shrink_to_fit,
                        indent=cell.alignment.indent,
                    )


def _save_workbook_safely(workbook, output_path: Path) -> Path:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    file_descriptor, temporary_name = tempfile.mkstemp(
        prefix=f'.{output_path.stem}_',
        suffix=output_path.suffix,
        dir=output_path.parent,
    )
    os.close(file_descriptor)
    temporary_path = Path(temporary_name)
    try:
        workbook.save(temporary_path)
        try:
            os.replace(temporary_path, output_path)
            return output_path
        except PermissionError as exc:
            if not output_path.exists():
                raise PermissionError(
                    f'无法写入 {output_path}，请检查文件夹权限。'
                ) from exc
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S_%f')
            versioned_path = output_path.with_name(
                f'{output_path.stem}_v{timestamp}{output_path.suffix}'
            )
            os.replace(temporary_path, versioned_path)
            return versioned_path
    finally:
        temporary_path.unlink(missing_ok=True)


def _sheet_for_document(workbook, document_type: str):
    target = 'COMMERCIAL INVOICE' if document_type == 'Commercial Invoice' else 'PACKING LIST'
    for sheet in workbook.worksheets:
        if target in sheet.title.upper():
            return sheet
    return None


def _number(values: Dict[str, Any], key: str) -> float:
    try:
        return float(values.get(key) or 0)
    except (TypeError, ValueError):
        return 0.0


def _fill_invoice_template(sheet, values: Dict[str, Any]) -> None:
    currency = str(values.get('currency', ''))
    gross_kg = _number(values, 'gross_weight')
    net_kg = _number(values, 'net_weight')
    gross_mt = gross_kg / 1000
    net_mt = net_kg / 1000
    cell_values = {
        'K3': f"COMMERCIAL INVOICE NO.\n{values.get('document_no', '')}",
        'A9': values.get('date', ''),
        'G9': values.get('order_no', ''),
        'K9': values.get('document_no', ''),
        'A12': values.get('payment_terms', ''),
        'G12': values.get('date', ''),
        'K12': values.get('lc_no', ''),
        'G15': values.get('destination_port', ''),
        'K15': values.get('lc_issue_date', ''),
        'A17': values.get('customer_name_en', ''),
        'A18': values.get('customer_address', ''),
        'A20': values.get('loading_port', ''),
        'A26': values.get('product_code') or values.get('product_name_en', ''),
        'B26': values.get('quantity', ''),
        'C26': gross_mt,
        'F26': values.get('product_description_en') or values.get('product_name_en', ''),
        'I26': net_mt,
        'L26': _number(values, 'unit_price'),
        'N26': _number(values, 'total_amount'),
        'A38': 'TOTAL',
        'B38': f"{values.get('quantity', '')} {values.get('unit', '')}",
        'C38': gross_mt,
        'I38': net_mt,
        'L38': values.get('incoterm', ''),
        'N38': _number(values, 'total_amount'),
        'A39': f"TOTAL AMOUNT: {currency} {_number(values, 'total_amount'):,.2f}",
        'A40': f"AMOUNT IN WORDS: {amount_in_english_words(values.get('total_amount'), currency)}",
    }
    for address, value in cell_values.items():
        _set_template_cell(sheet, address, value)
    for address in ('N26', 'N38'):
        cell = sheet[address]
        cell.number_format = '#,##0.00'
        cell.alignment = Alignment(
            horizontal='right',
            vertical='center',
            shrink_to_fit=True,
        )
    for address in ('E26', 'E38'):
        cell = sheet[address]
        cell.number_format = '#,##0.00'
        cell.alignment = Alignment(
            horizontal='right',
            vertical='center',
            shrink_to_fit=True,
        )


def _fill_packing_list_template(sheet, values: Dict[str, Any]) -> None:
    gross_kg = _number(values, 'gross_weight')
    net_kg = _number(values, 'net_weight')
    gross_mt = gross_kg / 1000
    net_mt = net_kg / 1000
    cell_values = {
        'F6': f"DATE: {values.get('date', '')}",
        'A10': values.get('customer_name_en', ''),
        'A11': values.get('customer_address', ''),
        'C13': f"NO. {values.get('document_no', '')}",
        'C14': values.get('order_no', ''),
        'A18': 1,
        'B18': values.get('product_name_en', ''),
        'C18': values.get('packaging_method', ''),
        'D18': net_mt,
        'E18': net_kg,
        'F18': gross_mt,
        'G18': gross_kg,
        'H18': values.get('product_specification', ''),
        'D22': net_mt,
        'E22': net_kg,
        'F22': gross_mt,
        'G22': gross_kg,
        'B24': values.get('product_name_en', ''),
        'B25': values.get('product_specification', ''),
        'B27': values.get('packaging_method', ''),
        'B29': values.get('hs_code', ''),
        'B30': f'{net_mt:g} MT',
        'B31': f'{net_kg:,.2f} KG',
        'B32': f'{gross_mt:g} MT',
        'B33': f'{gross_kg:,.2f} KG',
    }
    for address, value in cell_values.items():
        _set_template_cell(sheet, address, value)
    for row in range(18, 22):
        sheet.row_dimensions[row].height = 15
    sheet.row_dimensions[10].height = max(sheet.row_dimensions[10].height or 0, 24)
    sheet.row_dimensions[11].height = max(sheet.row_dimensions[11].height or 0, 42)
    for row in range(24, 43):
        label = sheet.cell(row, 1)
        value = sheet.cell(row, 2)
        has_value = value.value not in (None, '')
        if label.__class__.__name__ != 'MergedCell':
            label.alignment = Alignment(
                horizontal=label.alignment.horizontal,
                vertical='center',
                wrap_text=has_value,
            )
        if value.__class__.__name__ != 'MergedCell':
            value.alignment = Alignment(
                horizontal=value.alignment.horizontal,
                vertical='center',
                wrap_text=False,
            )
        text = str(label.value or '')
        column_width = sheet.column_dimensions['A'].width or 8.43
        characters_per_line = max(int(column_width * 0.75), 8)
        line_count = sum(
            max(1, (len(line) + characters_per_line - 1) // characters_per_line)
            for line in text.splitlines()
        ) if text and has_value else 1
        sheet.row_dimensions[row].height = max(18, line_count * 17)


def _apply_packing_list_table_borders(sheet) -> None:
    thin = Side(style='thin', color='000000')
    no_border = Side(style=None)
    for row in range(16, 23):
        for column in range(1, 10):
            cell = sheet.cell(row, column)
            cell.border = Border(
                left=no_border,
                right=thin if column in {1, 2, 3, 4, 5, 6, 7} else no_border,
                top=thin if row == 16 else no_border,
                bottom=thin if row in {17, 21, 22} else no_border,
            )


def _export_using_template(
    document_type: str,
    values: Dict[str, Any],
    output_path: Path,
    template_path: str,
) -> str:
    source = Path(template_path)
    if not source.exists():
        raise FileNotFoundError(f'未找到所选 Excel 模板：{source}')
    if source.suffix.lower() != '.xlsx':
        raise ValueError('Excel 单证模板必须是 .xlsx 文件。')

    workbook = load_workbook(source)
    invoice_sheet = _sheet_for_document(workbook, 'Commercial Invoice')
    packing_sheet = _sheet_for_document(workbook, 'Packing List')
    if invoice_sheet is None and packing_sheet is None:
        raise ValueError('所选 Excel 模板没有识别到 Commercial Invoice 或 Packing List 工作表。')

    if document_type == 'Commercial Invoice' and invoice_sheet is not None:
        _fill_invoice_template(invoice_sheet, values)
    elif document_type == 'Packing List' and packing_sheet is not None:
        _fill_packing_list_template(packing_sheet, values)
        _apply_packing_list_table_borders(packing_sheet)

    target_sheet = invoice_sheet if document_type == 'Commercial Invoice' else packing_sheet
    if document_type in {'Commercial Invoice', 'Packing List'} and target_sheet is None:
        workbook.close()
        raise ValueError(f'所选 Excel 模板中没有识别到 {document_type} 工作表。')
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = 'auto'
    target_sheet.page_setup.orientation = 'portrait'
    target_sheet.page_setup.paperSize = target_sheet.PAPERSIZE_A4
    target_sheet.page_setup.scale = None
    target_sheet.page_setup.fitToWidth = 1
    target_sheet.page_setup.fitToHeight = 1
    target_sheet.sheet_properties.pageSetUpPr.fitToPage = True
    _normalize_workbook_fonts(workbook)
    saved_path = _save_workbook_safely(workbook, output_path)
    workbook.close()
    return str(saved_path)


def _section(sheet, row: int, title: str, end_column: int = 5) -> int:
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=end_column)
    cell = sheet.cell(row, 1, title.upper())
    cell.fill = PatternFill('solid', fgColor=BLUE)
    cell.font = Font(name='Arial', size=10, bold=True, color=WHITE)
    cell.alignment = Alignment(vertical='center')
    sheet.row_dimensions[row].height = 20
    return row + 1


def _label_value(sheet, row: int, pairs: list[tuple[str, Any]]) -> int:
    column = 1
    for label, value in pairs:
        label_cell = sheet.cell(row, column, label)
        label_cell.fill = PatternFill('solid', fgColor=PALE_BLUE)
        label_cell.font = Font(name='Arial', size=9, bold=True, color=BLUE)
        label_cell.alignment = Alignment(vertical='top', wrap_text=True)
        value_cell = sheet.cell(row, column + 1, value)
        value_cell.font = Font(name='Arial', size=9, color='243447')
        value_cell.alignment = Alignment(vertical='top', wrap_text=True)
        column += 2
    sheet.row_dimensions[row].height = 30
    return row + 1


def _apply_table_borders(sheet, min_row: int, max_row: int, min_column: int, max_column: int) -> None:
    side = Side(style='thin', color=GRID)
    for row in sheet.iter_rows(min_row=min_row, max_row=max_row, min_col=min_column, max_col=max_column):
        for cell in row:
            cell.border = Border(left=side, right=side, top=side, bottom=side)


def export_document_excel(
    document_type: str,
    values: Dict[str, Any],
    output_path: str | None = None,
    template_path: str | None = None,
) -> str:
    target = Path(output_path) if output_path else settings.outputs_dir / values.get('order_no', 'document') / f'{document_type}_export.xlsx'
    target.parent.mkdir(parents=True, exist_ok=True)

    if template_path:
        template_output = _export_using_template(document_type, values, target, template_path)
        if template_output:
            return template_output

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = document_type[:31]
    sheet.sheet_view.showGridLines = False
    sheet.merge_cells('A1:F1')
    title = sheet['A1']
    title.value = str(values.get('company_name', ''))
    title.font = Font(name='Arial', size=16, bold=True, color=NAVY)
    title.alignment = Alignment(horizontal='center')
    sheet.merge_cells('A2:F2')
    title = sheet['A2']
    title.value = document_type.upper()
    title.font = Font(name='Arial', size=14, bold=True, color=BLUE)
    title.alignment = Alignment(horizontal='center')
    sheet.row_dimensions[1].height = 26
    sheet.row_dimensions[2].height = 24

    row = _section(sheet, 4, 'Document Information')
    row = _label_value(sheet, row, [
        ('Document No.', values.get('document_no', '')),
        ('Date', values.get('date', '')),
        ('Order No.', values.get('order_no', '')),
    ])
    row = _section(sheet, row, 'Parties')
    row = _label_value(sheet, row, [
        ('Seller', f"{values.get('company_name', '')}\n{values.get('company_address', '')}"),
        ('Buyer', f"{values.get('customer_name_en', '')}\n{values.get('customer_address', '')}"),
    ])
    row = _section(sheet, row, 'Shipment and Commercial Terms')
    row = _label_value(sheet, row, [
        ('Incoterm', values.get('incoterm', '')),
        ('Currency', values.get('currency', '')),
    ])
    row = _label_value(sheet, row, [
        ('Loading Port', values.get('loading_port', '')),
        ('Destination Port', values.get('destination_port', '')),
    ])
    row = _label_value(sheet, row, [
        ('Payment Terms', values.get('payment_terms', '')),
        ('Shipment Date', values.get('shipment_date', '')),
    ])

    row = _section(sheet, row, 'Goods and Amount')
    headers = ['No.', 'Description of Goods', 'Quantity', 'Unit Price', 'Amount', 'Currency']
    for column, value in enumerate(headers, start=1):
        cell = sheet.cell(row, column, value)
        cell.fill = PatternFill('solid', fgColor=NAVY)
        cell.font = Font(name='Arial', size=9, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    row += 1
    goods_row = row
    goods = [
        1,
        values.get('product_name_en', ''),
        f"{values.get('quantity', '')} {values.get('unit', '')}",
        values.get('unit_price', ''),
        values.get('total_amount', ''),
        values.get('currency', ''),
    ]
    for column, value in enumerate(goods, start=1):
        cell = sheet.cell(row, column, value)
        cell.font = Font(name='Arial', size=9, color='243447')
        cell.alignment = Alignment(vertical='top', wrap_text=True)
    row += 1
    if document_type == 'Packing List':
        for label, value in [
            ('Marks & Numbers', values.get('marks_and_numbers', '')),
            ('Net Weight (kg)', values.get('net_weight', '')),
            ('Gross Weight (kg)', values.get('gross_weight', '')),
            ('Measurement (CBM)', values.get('volume', '')),
        ]:
            sheet.cell(row, 1, label)
            sheet.cell(row, 2, value)
            sheet.cell(row, 1).font = Font(name='Arial', size=9, bold=True, color=BLUE)
            row += 1
    else:
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        sheet.cell(row, 1, 'TOTAL CONTRACT / INVOICE AMOUNT')
        sheet.cell(row, 5, values.get('total_amount', ''))
        sheet.cell(row, 6, values.get('currency', ''))
        sheet.cell(row, 1).font = Font(name='Arial', size=10, bold=True, color=NAVY)
        sheet.cell(row, 5).font = Font(name='Arial', size=10, bold=True, color=NAVY)
        row += 1
    _apply_table_borders(sheet, goods_row - 1, row - 1, 1, 6)

    row += 1
    row = _section(sheet, row, 'Packing and Remarks')
    row = _label_value(sheet, row, [
        ('Marks & Numbers', values.get('marks_and_numbers', '')),
        ('Remarks', values.get('remarks', '')),
    ])
    if values.get('lc_document_instruction'):
        row = _section(sheet, row, 'Letter of Credit Requirements')
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        sheet.cell(row, 1, values['lc_document_instruction'])
        sheet.cell(row, 1).alignment = Alignment(vertical='top', wrap_text=True)
        sheet.cell(row, 1).font = Font(name='Arial', size=9, color='243447')
        sheet.row_dimensions[row].height = max(30, 15 * str(values['lc_document_instruction']).count('\n') + 30)
        row += 1

    for column, width in enumerate([18, 34, 20, 18, 18, 15], start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = f'A{goods_row}'
    sheet.auto_filter.ref = f'A{goods_row - 1}:F{goods_row}'
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_margins = PageMargins(left=0.25, right=0.25, top=0.5, bottom=0.5, header=0.2, footer=0.2)
    sheet.print_title_rows = f'1:{goods_row - 1}'

    workbook.save(target)
    return str(target)


# 列表导出：客户、产品等成批数据的通用写出逻辑
FORMULA_PREFIXES = ('=', '+', '-', '@')


def sanitize_cell_value(value: Any) -> Any:
    """防止 Excel 公式注入：以 = + - @ 开头的文本前置英文单引号。"""
    if isinstance(value, str) and value.startswith(FORMULA_PREFIXES):
        return f"'{value}"
    return value


def dated_export_filename(prefix: str, moment: datetime | None = None) -> str:
    """导出文件名包含类型与导出日期，例如 客户_20261017.xlsx。"""
    when = moment or datetime.now()
    return f'{prefix}_{when.strftime("%Y%m%d")}.xlsx'


def _display_width(text: str) -> int:
    """中文按两个字符宽度估算列宽。"""
    return sum(2 if ord(character) > 127 else 1 for character in text)


def build_table_workbook(
    sheet_title: str,
    headers: list[str],
    rows: list[list[Any]],
    text_columns: set[int] | None = None,
    number_columns: set[int] | None = None,
) -> Workbook:
    """生成列表型工作簿：中文表头、文本列按文本保存、数值列带千分位。

    列序号从 1 开始，与 Excel 列号一致。
    """
    text_columns = text_columns or set()
    number_columns = number_columns or set()

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_title[:31]
    sheet.sheet_view.showGridLines = False

    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(1, column, str(header))
        cell.fill = PatternFill('solid', fgColor=NAVY)
        cell.font = Font(name='Arial', size=10, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    sheet.row_dimensions[1].height = 24

    for row_index, row in enumerate(rows, start=2):
        for column, raw_value in enumerate(row, start=1):
            value = raw_value if raw_value is not None else ''
            if column in text_columns and value != '':
                value = str(value)
            cell = sheet.cell(row_index, column, sanitize_cell_value(value))
            cell.font = Font(name='Arial', size=9, color='243447')
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            if column in text_columns:
                cell.number_format = '@'
            elif column in number_columns and isinstance(value, (int, float)):
                cell.number_format = '#,##0.00' if isinstance(value, float) else '#,##0'
                cell.alignment = Alignment(horizontal='right', vertical='top')

    last_row = max(len(rows) + 1, 1)
    _apply_table_borders(sheet, 1, last_row, 1, len(headers))
    for column, header in enumerate(headers, start=1):
        sample = rows[0][column - 1] if rows else ''
        width = max(_display_width(str(header)), _display_width(str(sample or '')))
        sheet.column_dimensions[get_column_letter(column)].width = min(max(width + 4, 10), 46)
    sheet.freeze_panes = 'A2'
    if rows:
        sheet.auto_filter.ref = f'A1:{get_column_letter(len(headers))}{last_row}'
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_margins = PageMargins(left=0.25, right=0.25, top=0.5, bottom=0.5, header=0.2, footer=0.2)
    sheet.print_title_rows = '1:1'
    return workbook


def workbook_bytes(workbook: Workbook) -> bytes:
    """把工作簿输出为内存字节，供下载使用，不在磁盘留文件。"""
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def build_table_export_file(
    prefix: str,
    sheet_title: str,
    headers: list[str],
    rows: list[list[Any]],
    text_columns: set[int] | None = None,
    number_columns: set[int] | None = None,
) -> tuple[str, bytes]:
    """一步生成可直接下载的（文件名, 字节流）。"""
    workbook = build_table_workbook(
        sheet_title, headers, rows, text_columns, number_columns
    )
    return dated_export_filename(prefix), workbook_bytes(workbook)
